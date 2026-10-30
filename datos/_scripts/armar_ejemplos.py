# Arma un Excel con unas 200 filas de ejemplo de cada fuente de la hoja ACC y un diccionario de sus campos, para que
# cada integrante vea qué trae cada fuente antes de analizarla.
# Uso: python armar_ejemplos.py   (usa los archivos ya bajados en datos/crudos)
import json
import random
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

DATOS = Path(__file__).resolve().parents[1]
CRUDOS, PROC = DATOS / "crudos", DATOS / "procesados"
SALIDA = DATOS / "ejemplos" / "Ejemplos de datos - LLIKA.xlsx"
N, SEMILLA = 200, 7


def ultimo(carpeta, patron="*.csv"):
    return sorted((CRUDOS / carpeta).glob(patron))[-1]


def vacios(df):
    return {c: float(df[c].isna().mean() + df[c].astype(str).str.strip().isin(["", "-", "None"]).mean() * df[c].notna().mean())
            for c in df.columns}


# (campo, qué es, para qué nos sirve). Oraciones completas, sin símbolos raros.
DIC = {
 "ACC-1": [
  ("entidad", "Es el nombre de la entidad pública que emitió la orden.", "Sirve para leer los casos; para cruzar se usa el RUC."),
  ("ruc_entidad", "Es el RUC de la entidad que emitió la orden.", "Es la llave para unir con el maestro de entidades (ACC-9) y con el scraper (ACC-2)."),
  ("departamento_entidad", "Es el departamento donde está la entidad.", "Permite separar por región."),
  ("tipoorden", "Indica si es orden de compra (bienes) u orden de servicio.", "Se usa para emparejar con el scraper."),
  ("nro_de_orden", "Es el número de la orden dentro de la entidad.", "Junto con la entidad y el tipo identifica la orden."),
  ("orden", "Es el código completo de la orden, con tipo, número, año y oficina.", "Identifica la orden y deja ver de qué año es."),
  ("descripcion_orden", "Es el texto libre con lo que se compró o contrató.", "En la H1′ sirve para juzgar si dos órdenes son del mismo objeto."),
  ("objetocontractual", "Indica si la orden es de bienes, de servicios o de consultoría de obras.", "Lectura secundaria de la H2, agrupando también por objeto."),
  ("estadocontratacion", "Es el estado de la orden: emitida, comprometida, devengada o anulada.", "Las anuladas deben excluirse antes de sumar."),
  ("tipodecontratacion", "Es el régimen de la contratación; las que empiezan con «Contrataciones hasta 8 UIT» son los contratos menores.", "Define el universo de la H2."),
  ("monto_total_orden_original", "Es el monto de la orden en su moneda original.", "Es lo que se suma contra el tope de 8 UIT."),
  ("moneda", "Es la moneda de la orden; casi todas están en soles.", "Las órdenes en otra moneda se excluyen en la H2."),
  ("ruc_contratista", "Es el RUC del proveedor; los que empiezan con 10 son personas naturales.", "Es la llave para unir con ACC-3, ACC-6 y ACC-7, y con el empleador previo de ACC-5."),
  ("nombre_razon_contratista", "Es el nombre o la razón social del proveedor.", "Sirve para leer los casos."),
  ("fecha_registro", "Es el día en que la entidad registró la orden en el SEACE; decide en qué archivo mensual aparece la orden.", "Muestra el atraso de registro, clave para el chequeo CAL-4."),
  ("fecha_de_emision", "Es el día en que se emitió la orden.", "Define el año de la orden en las tres hipótesis."),
  ("fecha_compromiso_presupuestal", "Es el día en que se comprometió el presupuesto de la orden.", "Ayuda a ver en qué etapa está la orden."),
  ("fecha_de_notificacion", "Es el día en que se notificó la orden al proveedor.", "Ayuda a ver en qué etapa está la orden."),
 ],
 "ACC-2": [
  ("N°", "Es el número de fila dentro del archivo que exportó el buscador.", "No se usa en el análisis."),
  ("Entidad", "Es el nombre de la entidad.", "Sirve para leer los casos."),
  ("Tipo de Orden", "Indica si es orden de compra (O/C) u orden de servicio (O/S).", "Se usa para emparejar con CONOSCE."),
  ("Número de Orden", "Es el número de la orden dentro de la entidad.", "Se usa para emparejar con CONOSCE."),
  ("Tipo de Contratación", "Es el régimen de la contratación, igual que en CONOSCE.", "Filtra los contratos menores."),
  ("Fecha de Emisión", "Es el día de emisión, con formato día/mes/año.", "Define el año de la orden."),
  ("Fecha de Compromiso", "Es el día del compromiso presupuestal; suele venir vacío.", "Uso secundario."),
  ("Monto", "Es el monto como texto, con el símbolo de la moneda y comas de miles (por ejemplo «S/. 1,159.11»).", "Hay que limpiarlo antes de sumar."),
  ("RUC", "Es el RUC del proveedor.", "Es la llave con las demás fuentes."),
  ("Denominación o Razón Social", "Es el nombre del proveedor.", "Sirve para leer los casos."),
  ("Estado", "Es el estado de la orden, por ejemplo devengada.", "Permite excluir las anuladas."),
  ("Estado de Registro", "Indica si la orden se registró dentro o fuera del plazo.", "Confirma el atraso de registro que se ve en CONOSCE."),
  ("Observaciones", "Son notas del registro; casi siempre está vacío.", "No se usa en el análisis."),
  ("ruc_entidad (de la carpeta)", "No viene en el archivo: sale del nombre de la carpeta donde el scraper lo guardó.", "Es la única forma de saber el RUC de la entidad en esta fuente."),
  ("mes (del archivo)", "No viene en el archivo: es el mes consultado, tomado del nombre del archivo.", "Ubica la orden en el tiempo."),
 ],
 "ACC-3": [
  ("FECHA_CORTE", "Es la fecha de corte del archivo (año, mes y día); es la misma en todas las filas.",
   "Es una foto mensual: el corte es el día 1 y se publica entre el 5 y el 7. La wiki del OECE guarda una versión por mes desde noviembre de 2021."),
  ("TIPO_DOCUMENTO", "Es el tipo de documento que se declaró para el integrante (DNI, carné de extranjería, RUC u otro). No es confiable: 1 585 filas dicen DNI y traen un RUC.",
   "Para saber si el integrante es persona o empresa se mira el número, no este campo."),
  ("NUMERO_DOCUMENTO", "Es el número de documento; algunos vienen con un tabulador pegado al inicio.", "Une a las empresas que comparten socios; hay que limpiarlo antes de cruzar."),
  ("NOMBRE_RAZONODENOMINACIONSOCIAL", "Es el nombre de la persona (o de la empresa) que integra al proveedor, no el del proveedor.",
   "Sirve para leer los casos. La razón social del proveedor no está en este archivo: sale del padrón de SUNAT (ACC-6) por el RUC."),
  ("RUC", "Es el RUC del proveedor; siempre es una persona jurídica (empresas, asociaciones, cooperativas). Los códigos que empiezan con 99 u 88 son proveedores extranjeros sin RUC peruano.",
   "Se une con ruc_contratista de ACC-1."),
  ("TIPO_CONF_JURIDICA", "Es el rol de la persona: socio o accionista, órgano de administración o representante legal. Una misma persona puede tener varios roles en el mismo proveedor, una fila por rol.",
   "Permite armar la red de proveedores vinculados; para el representante legal se filtra REPRESENTANTE_LEGAL."),
  ("FECHA_INICIO_VIGENCIA", "Es una sola fecha por proveedor, igual en todas sus filas: la de su último trámite en el RNP (inscripción, renovación o actualización). Cambia cada vez que hace un trámite. No es la fundación de la empresa ni el inicio del cargo de cada persona.",
   "No sirve como antigüedad ni como fecha del cargo. Para saber desde cuándo existe un vínculo se comparan los cortes mensuales de la wiki del OECE."),
  ("ID_FORMA_SOCIETARIA", "Es el código de la forma societaria.", "Uso secundario."),
  ("DE_FORMA_SOCIETARIA", "Es la forma societaria, por ejemplo SAC o EIRL.", "Uso secundario."),
 ],
 "ACC-4": [
  ("CODIGO_DDJJ", "Es el código de la declaración jurada.", "Es la llave para unir con los empleos previos (ACC-5)."),
  ("TIPO_DOCUMENTO_FUNCIONARIO", "Es solo el tipo de documento del funcionario. El número no se publica en ninguna versión abierta (tampoco en los JSON de 2022).",
   "El DNI se puede recuperar por el nombre completo en el padrón de SUNAT: el 74,3 % de los declarantes tiene un solo RUC 10 con su nombre."),
  ("APELLIDO_PATERNO, APELLIDO_MATERNO y NOMBRES", "Son los nombres del funcionario. Están en el archivo original; en este ejemplo los quitamos por ser datos personales.",
   "Son la única forma de identificar al declarante: sirven para cruzar con los socios y representantes de ACC-3 y para recuperar su DNI en el padrón de SUNAT."),
  ("ENTIDAD", "Es el nombre de la entidad donde trabaja el funcionario; no trae RUC. Solo el 56 % de los nombres calza exacto con el maestro de entidades.",
   "En la H3′ se empareja por nombre con el maestro de entidades (ACC-9); hace falta quitar tildes y sufijos como « - CGR»."),
  ("CARGO", "Es el cargo del funcionario, en texto libre. En el 90 % de las filas es genérico («OTRO CARGO FUNCION O LABOR» o «COLABORADOR»).",
   "Permite enfocarse en cargos que deciden compras (comité de selección, logística, gerencia) cuando el texto lo dice."),
  ("FECHA_ELABORACION_DDJJ", "Es la fecha en que se elaboró la declaración, con año de dos cifras.", "Uso secundario."),
  ("FECHA_PRESENTACION_DDJJ", "Es la fecha en que se presentó la declaración, con año de dos cifras. Van de marzo de 2022 a julio de 2026.", "Ubica la declaración en el tiempo."),
  ("FECHA_INICIO_CARGO", "Es la fecha en que el funcionario empezó en el cargo, con año de dos cifras: 12 fechas salen en 2050 o 2060 y son de 1950 o 1960.",
   "En la H3′ define si el cargo estaba vigente en 2024."),
  ("FECHA_CESE_CARGO", "Es la fecha en que dejó el cargo; queda vacía si sigue en él (61 % de las filas).", "En la H3′ define si el cargo estaba vigente en 2024."),
 ],
 "ACC-5": [
  ("CODIGO_DDJJ", "Es el código de la declaración jurada a la que pertenece este empleo. Un empleo se repite en cada declaración de la misma persona.",
   "Une con ACC-4, que pone el nombre y la entidad actual; ACC-5 no trae ni nombre ni documento del funcionario."),
  ("RUC_ENTIDAD_LABORO", "Es el identificador del empleador anterior, con un prefijo («R.U.C.:», «RUC:», «D.N.I.:»). El 98,1 % trae 11 dígitos; el 1,6 % es un DNI y el 0,3 % viene vacío.",
   "En la H3′ es el RUC que se busca entre los proveedores. El 86 % de los empleos previos es en una entidad pública."),
  ("NOMBRE_ENTIDAD_LABORO", "Es el nombre del empleador anterior.", "Sirve para leer los casos."),
  ("CARGO_ENTIDAD_LABORO", "Es el cargo que tuvo el funcionario en ese empleo, en texto libre.", "Permite ver, por ejemplo, quién fue especialista en contrataciones."),
  ("FECHA_INICIO_CARGO", "Es la fecha de inicio de ese empleo, con formato día/mes/año de cuatro cifras.", "Ubica el empleo previo en el tiempo."),
  ("FECHA_CESE_CARGO", "Es la fecha de fin de ese empleo; vacía en el 30 % de las filas.", "Ubica el empleo previo en el tiempo."),
  ("FECHA_ELABORACION_DDJJ", "Es la fecha de elaboración de la declaración.", "Uso secundario."),
  ("FECHA_PRESENTACION_DDJJ", "Es la fecha de presentación de la declaración.", "Uso secundario."),
 ],
 "ACC-6": [
  ("RUC", "Es el RUC del contribuyente.", "Se une con los RUC de las demás fuentes."),
  ("NOMBRE O RAZÓN SOCIAL", "Es el nombre de la persona o la razón social de la empresa.", "Sirve para leer los casos."),
  ("ESTADO DEL CONTRIBUYENTE", "Indica si el RUC está activo o dado de baja, y por qué.", "Marca proveedores dados de baja que siguen recibiendo órdenes."),
  ("CONDICIÓN DE DOMICILIO", "Indica si SUNAT ubicó el domicilio (habido) o no (no habido).", "Un proveedor no habido es una señal de riesgo."),
  ("UBIGEO", "Es el código de ubicación del domicilio fiscal.", "Permite ver si el proveedor es de la misma zona que la entidad."),
  ("TIPO DE VÍA, NOMBRE DE VÍA, CÓDIGO DE ZONA, TIPO DE ZONA, NÚMERO, INTERIOR, LOTE, MANZANA y KILÓMETRO", "Son las partes de la dirección del domicilio fiscal; SUNAT pone un guion cuando no hay dato.", "Permiten detectar proveedores que comparten dirección."),
  ("DEPARTAMENTO", "Es el número de departamento de la dirección (el piso o la unidad), no la región.", "Parte de la dirección; no confundir con la región."),
 ],
 "ACC-7": [
  ("FECHA_CORTE", "Es la fecha de corte del archivo.", "Recuerda que es una foto."),
  ("RUC", "Es el RUC del proveedor sancionado.", "Se une con ruc_contratista de ACC-1."),
  ("NOMBRE_RAZONODENOMINACIONSOCIAL", "Es el nombre del proveedor sancionado.", "Sirve para leer los casos."),
  ("FECHA_INICIO", "Es el día en que empieza la sanción.", "Permite saber si la sanción es anterior o posterior a una orden."),
  ("FECHA_FIN", "Es el día en que termina la sanción; queda vacío si es definitiva.", "Permite saber si la sanción estaba vigente."),
  ("NUMERO_RESOLUCION", "Es el número de la resolución del Tribunal de Contrataciones.", "Sirve para citar el caso."),
  ("ID_MOTIVO_INFRACCION", "Es el código del motivo de la sanción.", "Uso secundario."),
  ("DE_MOTIVO_INFRACCION", "Es el motivo de la sanción, por ejemplo presentar información falsa.", "Permite separar tipos de infracción."),
 ],
 "ACC-8": [
  ("NumeroInforme", "Es el número oficial del informe de control.", "Sirve para citar el informe."),
  ("CodigoInforme", "Es el código interno del informe en el buscador.", "Identifica el informe."),
  ("Descripcion", "Es el título del informe, que suele decir qué se fraccionó y en qué periodo.", "Permite saber la materia y el periodo del caso."),
  ("Entidad", "Es la entidad auditada; viene sin RUC.", "Se empareja por nombre con el maestro de entidades (ACC-9)."),
  ("Departamento, Provincia y Distrito", "Son la ubicación de la entidad auditada.", "Permite separar por región."),
  ("Sector y NivelGobierno", "Son el sector y el nivel de gobierno de la entidad.", "Uso secundario."),
  ("ModalidadServicio", "Es la modalidad del servicio de control, por ejemplo servicio de control específico.", "Indica si el servicio podía determinar responsabilidades."),
  ("ServicioControl", "Indica si el control fue posterior, simultáneo u otro.", "Uso secundario."),
  ("Responsabilidad y EsConResponsabilidad", "Indican si el informe identificó responsabilidades y de qué tipo.", "Separa los informes que llegaron a responsabilidad."),
  ("TotalFuncionarios", "Es el número de funcionarios comprendidos; los nombres no se publican.", "Uso secundario."),
  ("Periodo y Aniofechafinejecucion", "Son el año del informe y el año en que terminó el servicio.", "Ubica el informe en el tiempo."),
  ("FechaEmision, FechaPublicacion y FechaFinEjecucion", "Son las fechas de emisión, publicación y fin del servicio de control.", "Son posteriores a las compras: no pueden usarse como variable del modelo."),
  ("ResumenEjecutivo y RutaCloudInforme", "Son los enlaces a la ficha resumen y al informe en PDF.", "Permiten leer el caso completo."),
  ("TieneApendice, TieneAnexos y TotalRecomendaciones", "Indican si el informe trae apéndices o anexos y cuántas recomendaciones tiene.", "Uso secundario."),
  ("Otros campos (Correlativo, TotalRows, TotalPages, Text, Publico, EsCovid, rutas y códigos vacíos)", "Son datos técnicos del buscador o vienen vacíos.", "No aportan al análisis."),
 ],
 "ACC-9": [
  ("RUC", "Es el RUC de la entidad contratante.", "Se une con ruc_entidad de ACC-1."),
  ("NOMBRE_DE_ENTIDAD", "Es el nombre oficial de la entidad.", "Sirve para emparejar por nombre las declaraciones (ACC-4) y los informes (ACC-8)."),
  ("DEPARTAMENTO, PROVINCIA y DISTRITO", "Son la ubicación de la entidad.", "Permite separar por región; en la H3′ las permutaciones se hacen dentro del departamento."),
  ("CODIGO_SIAF", "Es el código de la unidad ejecutora en el SIAF del MEF.", "Permite unir con datos de presupuesto si hiciera falta."),
  ("CODCONSUCODE", "Es el código interno de la entidad en el OECE.", "Uso secundario."),
  ("ESTADO", "Indica si la entidad está activa o inactiva.", "Uso secundario."),
  ("ULTIMAACTUALIZACION", "Es la fecha de la última actualización del registro, con año de dos cifras.", "Uso secundario."),
 ],
}

ACC8_CAMPOS = ["NumeroInforme", "CodigoInforme", "Descripcion", "Entidad", "Departamento", "Provincia", "Distrito",
               "Sector", "NivelGobierno", "ModalidadServicio", "ServicioControl", "Responsabilidad", "EsConResponsabilidad",
               "TotalFuncionarios", "Periodo", "Aniofechafinejecucion", "FechaEmision", "FechaPublicacion",
               "FechaFinEjecucion", "ResumenEjecutivo", "RutaCloudInforme", "TieneApendice", "TieneAnexos",
               "TotalRecomendaciones"]

if __name__ == "__main__":
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    hojas, completo, vac = {}, {}, {}

    # ACC-1: un mes reciente y completo de CONOSCE (agosto de 2026)
    o = pd.read_parquet(PROC / "acc1_conosce" / "ordenes_2026-08.parquet").drop(columns=["archivo"])
    vac["ACC-1"], completo["ACC-1"] = vacios(o), len(o)
    hojas["ACC-1"] = o.sample(N, random_state=SEMILLA).sort_values(["entidad", "nro_de_orden"])
    proveedores, entidades = set(hojas["ACC-1"]["ruc_contratista"]), set(hojas["ACC-1"]["ruc_entidad"])

    # ACC-2: archivos del scraper de Cusco 2024 (40 al azar)
    xls = sorted((CRUDOS / "acc2_scraper" / "repo" / "data" / "raw" / "cusco").glob("*/2024/*.xls"))
    partes = []
    for f in random.Random(SEMILLA).sample(xls, 40):
        df = pd.read_excel(f, engine="calamine", dtype=str)
        df["ruc_entidad (de la carpeta)"], df["mes (del archivo)"] = f.parts[-3], f"2024-{f.stem}"
        partes.append(df)
    s = pd.concat(partes, ignore_index=True)
    vac["ACC-2"], completo["ACC-2"] = vacios(s), None
    hojas["ACC-2"] = s.sample(min(N, len(s)), random_state=SEMILLA)

    # ACC-3: conformación jurídica, primero los proveedores del ejemplo de ACC-1
    c = pd.read_csv(ultimo("acc3_conformacion"), sep="|", dtype=str, encoding="cp1252", encoding_errors="replace",
                    on_bad_lines="skip")
    vac["ACC-3"], completo["ACC-3"] = vacios(c), len(c)
    ligadas = c[c["RUC"].isin(proveedores)]
    hojas["ACC-3"] = pd.concat([ligadas, c.drop(ligadas.index).sample(max(0, N - len(ligadas)), random_state=SEMILLA)]).head(N)

    # ACC-4 y ACC-5: declaraciones al azar y sus empleos previos (mismo CODIGO_DDJJ)
    d4 = pd.read_csv(ultimo("acc4_dji_declaraciones"), dtype=str, encoding="utf-8")
    vac["ACC-4"], completo["ACC-4"] = vacios(d4), len(d4)
    hojas["ACC-4"] = d4.sample(N, random_state=SEMILLA).drop(columns=["APELLIDO_PATERNO", "APELLIDO_MATERNO", "NOMBRES"])
    d5 = pd.read_csv(ultimo("acc5_dji_empleos"), dtype=str, encoding="utf-8")
    vac["ACC-5"], completo["ACC-5"] = vacios(d5), len(d5)
    hojas["ACC-5"] = d5[d5["CODIGO_DDJJ"].isin(set(hojas["ACC-4"]["CODIGO_DDJJ"]))].head(N)

    # ACC-6: padrón de SUNAT, buscando los RUC del ejemplo de ACC-1 (se lee por partes: pesa 1,6 GB)
    filas, total, vacias = [], 0, None
    with zipfile.ZipFile(ultimo("acc6_padron_sunat", "*.zip")) as z, z.open("padron_reducido_ruc.txt") as fh:
        for t in pd.read_csv(fh, sep="|", dtype=str, encoding="latin-1", usecols=range(15), chunksize=2_000_000,
                             quoting=3, on_bad_lines="skip"):
            total += len(t)
            v = t.isin(["-", ""]).sum() + t.isna().sum()
            vacias = v if vacias is None else vacias + v
            filas.append(t[t["RUC"].isin(proveedores | entidades)])
    pad = pd.concat(filas, ignore_index=True)
    vac["ACC-6"], completo["ACC-6"] = (vacias / total).to_dict(), total
    hojas["ACC-6"] = pad.head(N)

    # ACC-7: sancionados al azar, más los del ejemplo de ACC-1 si los hubiera
    sa = pd.read_csv(ultimo("acc7_sancionados"), sep="|", dtype=str, encoding="cp1252", encoding_errors="replace")
    vac["ACC-7"], completo["ACC-7"] = vacios(sa), len(sa)
    hojas["ACC-7"] = pd.concat([sa[sa["RUC"].isin(proveedores)], sa.sample(N, random_state=SEMILLA)]).drop_duplicates().head(N)

    # ACC-8: los 176 informes, solo con los campos que aportan
    m = pd.DataFrame(json.load(open(CRUDOS / "acc8_informes_cgr" / "cgr-buscador-informes-fraccionamiento-176-metadatos.json",
                                    encoding="utf-8")))[ACC8_CAMPOS]
    vac["ACC-8"], completo["ACC-8"] = vacios(m), len(m)
    hojas["ACC-8"] = m

    # ACC-9: maestro de entidades, primero las entidades del ejemplo de ACC-1
    en = pd.read_csv(ultimo("acc9_entidades"), sep="|", dtype=str, encoding="cp1252", encoding_errors="replace")
    vac["ACC-9"], completo["ACC-9"] = vacios(en), len(en)
    ligadas = en[en["RUC"].isin(entidades)]
    hojas["ACC-9"] = pd.concat([ligadas, en.drop(ligadas.index).sample(N - len(ligadas), random_state=SEMILLA)])

    # Índice y diccionario
    NOMBRE = {"ACC-1": "ACC-1 Órdenes CONOSCE", "ACC-2": "ACC-2 Scraper SEACE", "ACC-3": "ACC-3 Conformación",
              "ACC-4": "ACC-4 Declaraciones", "ACC-5": "ACC-5 Empleos previos", "ACC-6": "ACC-6 Padrón SUNAT",
              "ACC-7": "ACC-7 Sancionados", "ACC-8": "ACC-8 Informes CGR", "ACC-9": "ACC-9 Entidades"}
    M = {k: f"{v:,}".replace(",", " ") for k, v in completo.items() if v}  # miles separados con espacio
    indice = pd.DataFrame([
     ["ACC-1", "Órdenes de compra y de servicio (CONOSCE)", "OECE", "Una orden de compra o de servicio.",
      f"Unas {M['ACC-1']} filas por mes; el ejemplo sale de agosto de 2026.", "2018 a 2026",
      "Por ruc_entidad con ACC-9 y ACC-2; por ruc_contratista con ACC-3, ACC-6 y ACC-7.", "H1′, H2 y H3′"],
     ["ACC-2", "Órdenes del buscador del SEACE que bajó el scraper del equipo", "OECE (copia en el repositorio del equipo)",
      "Una orden; los archivos vienen por entidad y por mes.", "277 071 archivos mensuales; el ejemplo sale de 40 archivos de Cusco 2024.",
      "2013 a marzo de 2026", "Por RUC de la entidad, tipo y número de orden con ACC-1.", "H1′ (es lo que hoy ordena nexus-core)"],
     ["ACC-3", "Conformación jurídica de los proveedores", "OECE", "Una persona con un rol dentro de un proveedor.",
      f"{M['ACC-3']} filas.", "Foto mensual; la wiki del OECE guarda una versión por mes desde 11/2021", "Por RUC con ruc_contratista de ACC-1.", "Señal de red (vínculo societario)"],
     ["ACC-4", "Declaraciones juradas de intereses: listado de declaraciones", "Contraloría", "Una declaración jurada.",
      f"{M['ACC-4']} filas.", "Marzo de 2022 a julio de 2026", "Por CODIGO_DDJJ con ACC-5; por nombre de entidad con ACC-9.", "H3′"],
     ["ACC-5", "Declaraciones juradas de intereses: empleos previos", "Contraloría", "Un empleo previo declarado.",
      f"{M['ACC-5']} filas.", "Según la fecha de la declaración", "Por CODIGO_DDJJ con ACC-4; por el RUC del empleador con ACC-1.", "H3′"],
     ["ACC-6", "Padrón reducido de contribuyentes", "SUNAT", "Un RUC.", f"{M['ACC-6']} filas.",
      "Foto del día de descarga", "Por RUC con todas las demás.", "H3′ y señales de riesgo del proveedor"],
     ["ACC-7", "Proveedores sancionados con inhabilitación", "OECE (Tribunal de Contrataciones)", "Una sanción.",
      f"{M['ACC-7']} filas.", "Sanciones desde 1993", "Por RUC con ruc_contratista de ACC-1.", "Contexto; no como variable del modelo"],
     ["ACC-8", "Informes de control por fraccionamiento", "Contraloría", "Un informe de control.",
      f"{M['ACC-8']} fichas.", "Informes de 2017 a 2026", "Por nombre de entidad con ACC-9.", "Repetición de la prueba 35 contra 35"],
     ["ACC-9", "Maestro de entidades contratantes", "OECE", "Una entidad contratante.", f"{M['ACC-9']} filas.",
      "Foto del día de descarga", "Por RUC con ruc_entidad de ACC-1.", "H3′ y emparejamiento de nombres"],
    ], columns=["ID", "Fuente", "Quién la publica", "Qué es una fila", "Tamaño de la fuente completa", "Periodo",
                "Se une con otras fuentes por", "La usa"])
    indice["Hoja con ejemplos"] = indice["ID"].map(NOMBRE)
    indice["Cómo se lee el archivo original"] = indice["ID"].map({
        "ACC-1": "Un Excel por mes (.xlsx) con una hoja llamada CONOSCE.",
        "ACC-2": "Un Excel antiguo (.xls) por entidad y por mes; el monto viene como texto.",
        "ACC-3": "Texto separado por barras verticales (|), con codificación de Windows (cp1252).",
        "ACC-4": "CSV separado por comas, en UTF-8; las fechas traen el año con dos cifras.",
        "ACC-5": "CSV separado por comas, en UTF-8.",
        "ACC-6": "Texto separado por barras verticales (|) dentro de un .zip, con codificación Latin-1; un guion significa que no hay dato.",
        "ACC-7": "Texto separado por barras verticales (|), con codificación de Windows (cp1252).",
        "ACC-8": "JSON bajado del buscador de informes de la Contraloría.",
        "ACC-9": "Texto separado por barras verticales (|), con codificación de Windows (cp1252).",
    })

    def legible(v):
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        if hasattr(v, "strftime"):
            return v.strftime("%Y-%m-%d")
        return str(v).strip()

    dic = []
    for acc, campos in DIC.items():
        h = hojas[acc]
        # todos los ejemplos de una fuente salen de la misma fila: la más completa
        llenos = h.apply(lambda col: ~col.astype(str).str.strip().isin(["", "-", "None", "nan", "NaT"]) & col.notna())
        modelo = h.loc[llenos.sum(axis=1).idxmax()]
        for campo, que, uso in campos:
            nombres = [x.strip() for x in campo.replace(" y ", ", ").split(",")]
            vs = [vac[acc].get(x) for x in nombres if x in vac[acc]]
            ej = ""
            if nombres[0] in h:
                ej = legible(modelo[nombres[0]]) if llenos.loc[modelo.name, nombres[0]] else ""
                if not ej:  # si la fila modelo no lo trae, el primer valor no vacío de la columna
                    resto = h[nombres[0]][llenos[nombres[0]]]
                    ej = legible(resto.iloc[0]) if len(resto) else ""
            # en campos agrupados se informa el del campo mejor llenado
            dic.append([acc, campo, que, ej[:80], round(min(vs), 3) if vs else None, uso])
    dicc = pd.DataFrame(dic, columns=["Fuente", "Campo", "Qué es", "Ejemplo", "Vacío en la fuente", "Para qué nos sirve"])

    with pd.ExcelWriter(SALIDA, engine="openpyxl", datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as w:
        indice.to_excel(w, sheet_name="Índice", index=False)
        dicc.to_excel(w, sheet_name="Diccionario", index=False)
        for acc, df in hojas.items():
            df.to_excel(w, sheet_name=NOMBRE[acc], index=False)

    # Formato sobrio: encabezado en negrita, fila fija, filtros y anchos legibles; sin colores
    wb = load_workbook(SALIDA)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(name="Calibri", bold=True)
        texto_largo = ws.title in ("Índice", "Diccionario")
        for col in ws.columns:
            largo = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, largo + 2), 55 if texto_largo else 40)
            if texto_largo:
                for c in col:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in wb["Diccionario"]["E"][1:]:
        c.number_format = "0.0%"
    wb.save(SALIDA)
    print("listo:", SALIDA)
    for acc, df in hojas.items():
        print(f"  {NOMBRE[acc]:24} {len(df):>4} filas de ejemplo, {df.shape[1]:>2} campos")
